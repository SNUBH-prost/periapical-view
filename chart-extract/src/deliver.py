"""
전체 결과 파일 생성 (노트 추출 + HbA1c 조인 + 연구시트 양식으로 출력).

사용법:
    python -m src.deliver --raw 원본노트.xlsx --hba1c HbA1c.xlsx \
                          --template 연구시트.xlsx --out 결과.xlsx

- --raw      : 외래경과 노트 엑셀 (한 줄=한 노트). 필수.
- --hba1c    : HbA1c 검사결과 엑셀 (환자번호+접수일시+수치). 있으면 식립일 기준
               직전(술전)/직후(술후)로 조인. 없으면 HbA1c 칸은 비움.
- --template : 연구시트 엑셀. 있으면 그 첫 줄(열 제목)을 그대로 사용해 채운다.
               없으면 기본 열을 쓴다.
- 당뇨(DM)·투약·결과(골흡수/합병증/fail) 열은 다른 소스라서 '빈칸'으로 둔다.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook

from . import excel_io
from .main import process

DEFAULT_HEADERS = [
    "Serial No.", "환자번호", "나이", "생년월일", "성별",
    "임플란트 식립 부위(상악/하악)", "임플란트 식립 시기", "임플란트 회사",
    "임플란트 직경", "임플란트길이", "골이식 여부",
    "당화혈색소(수술전)", "당화혈색소 측정시기(수술전)",
    "당화혈색소(수술후)", "당화혈색소 측정시기(수술후)", "당화혈색소 (술전- 술후)",
    "ISQ 측정량(협)", "ISQ 측정량(설)", "ISQ 측정평균",
]
HELP_HEADERS = ["치식", "ISQ원본", "ISQ측정일"]


def load_hba1c(path):
    """{환자번호: [(날짜, 값float or None, 값원본), ...] 정렬됨}"""
    if not path:
        return {}
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hb = {}
    for r in rows[1:]:
        p = str(r[1]); d = str(r[6])[:10]; v = r[5]
        try:
            vf = float(v)
        except (TypeError, ValueError):
            vf = None
        hb.setdefault(p, []).append((d, vf, v))
    for p in hb:
        hb[p].sort()
    return hb


def join_hba1c(hb, patient, surgery_date):
    """식립일 직전=술전, 직후=술후."""
    s = hb.get(patient, [])
    pre = [x for x in s if surgery_date and x[0] <= surgery_date]
    post = [x for x in s if surgery_date and x[0] > surgery_date]
    return (pre[-1] if pre else None), (post[0] if post else None)


def build(raw, hba1c_path, template, out):
    records, meta = excel_io.read_notes(raw)
    rows = process(records)

    demo = {}
    for r in records:
        p = str(r["patient"])
        demo.setdefault(p, (r["age"], r["birth"], r["sex"]))

    hb = load_hba1c(hba1c_path)

    if template:
        HDR = [c.value for c in load_workbook(template, data_only=True).active[1]]
        HDR = [h for h in HDR if h is not None]
    else:
        HDR = list(DEFAULT_HEADERS)

    def col(name):
        for i, h in enumerate(HDR):
            if h and str(h).strip() == name.strip():
                return i
        return None

    wb = Workbook(); ws = wb.active
    ws.append(HDR + HELP_HEADERS)
    serial = 0
    n_isq = n_hb = 0
    for r in rows:
        pat = str(r[1]); tooth = r[17]; sdate = r[5]
        age, birth, sex = demo.get(pat, ("", "", ""))
        pre, post = join_hba1c(hb, pat, sdate)
        line = [None] * len(HDR)

        def put(name, val):
            i = col(name)
            if i is not None:
                line[i] = val

        serial += 1
        put("Serial No.", serial)
        put("환자번호", pat)
        put("나이", age)
        put("생년월일", str(birth)[:10] if birth else "")
        put("성별", sex)
        put("임플란트 식립 부위(상악/하악)", f"{tooth}i")
        put("임플란트 식립 시기", sdate)
        put("임플란트 회사", r[6])
        put("임플란트 직경", r[7])
        put("임플란트길이", r[8])
        put("골이식 여부", r[9])
        if pre:
            put("당화혈색소(수술전)", pre[2]); put("당화혈색소 측정시기(수술전)", pre[0]); n_hb += 1
        if post:
            put("당화혈색소(수술후)", post[2]); put("당화혈색소 측정시기(수술후)", post[0])
        if pre and post and pre[1] is not None and post[1] is not None:
            put("딩화혈색소 (술전- 술후)", round(pre[1] - post[1], 1))
            put("당화혈색소 (술전- 술후)", round(pre[1] - post[1], 1))
        put("ISQ 측정량(협)", r[14]); put("ISQ 측정량(설)", r[15]); put("ISQ 측정평균", r[16])
        if r[14] != "":
            n_isq += 1
        line += [tooth, r[18], r[19]]
        ws.append(line)

    wb.save(out)
    print(f"저장 → {out}")
    print(f"임플란트 {serial}건 | ISQ 채움 {n_isq} | HbA1c(술전) 채움 {n_hb}")
    print("※ 당뇨·투약·결과(골흡수/합병증/fail) 열은 다른 소스라 빈칸입니다.")


def main():
    ap = argparse.ArgumentParser(description="노트 추출 + HbA1c 조인 → 연구시트 결과 생성")
    ap.add_argument("--raw", required=True)
    ap.add_argument("--hba1c", default=None)
    ap.add_argument("--template", default=None)
    ap.add_argument("--out", default="추출결과.xlsx")
    a = ap.parse_args()
    for f in (a.raw, a.hba1c, a.template):
        if f and not Path(f).exists():
            print(f"[오류] 파일 없음: {f}"); return 1
    build(a.raw, a.hba1c, a.template, a.out)
    return 0


if __name__ == "__main__":
    main()
